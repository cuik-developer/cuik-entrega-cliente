"""
Cuik-branded PowerPoint helper (python-pptx).

    from cuik_pptx import CuikDeck
    d = CuikDeck()
    d.cover("Tu programa de puntos", "Guía para Nombre del negocio", tag="Setiembre 2026")
    d.section("Cómo suman puntos tus clientes")
    d.bullets("Así funciona", ["1 punto por cada S/ 1", "El pase se actualiza solo", "..."])
    d.two_columns("Antes y ahora", ["...", "..."], ["...", "..."], left_title="Antes", right_title="Ahora")
    d.stats("Tu programa hoy", [("6", "clientes"), ("30", "visitas"), ("120", "puntos en circulación")])
    d.table("Catálogo de premios", ["Premio", "Puntos"], [["Café", "50"], ["Postre", "80"]])
    d.closing("¿Dudas?", "francesco.leon@cuik.org · cuik.org")
    d.save("salida.pptx")

Fonts are referenced by name (Poppins). Install assets/fonts/*.ttf on the
machine that opens/exports the file for full fidelity; otherwise Arial falls in.
"""

from __future__ import annotations

import os
from typing import Iterable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(os.path.dirname(HERE), "assets")
LOGO_MARK = os.path.join(ASSETS, "logo", "logo-mark.jpeg")
LOGO_MARK_ON_BLUE = os.path.join(ASSETS, "logo", "logo-mark-on-blue.png")
WORDMARK = os.path.join(ASSETS, "logo", "wordmark-blue-on-white.png")
ASTERISK = {
    "orange": os.path.join(ASSETS, "logo", "asterisk-orange.png"),
    "blue": os.path.join(ASSETS, "logo", "asterisk-blue.png"),
    "black": os.path.join(ASSETS, "logo", "asterisk-black.png"),
    "outline": os.path.join(ASSETS, "logo", "asterisk-outline.png"),
}

# ── Tokens ──────────────────────────────────────────────────────────
BLUE = RGBColor(0x0E, 0x70, 0xDB)
ORANGE = RGBColor(0xFF, 0x48, 0x10)
INK = RGBColor(0x23, 0x1F, 0x20)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xF8, 0xF8, 0xF8)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LINE = RGBColor(0xE5, 0xE7, 0xEB)

FONT_DISPLAY = "Poppins Black"  # Cocogoose only for the wordmark image
FONT_BOLD = "Poppins"  # bold=True
FONT_BODY = "Poppins"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.7)


def _font(run, size: int, color: RGBColor = INK, bold: bool = False, black: bool = False):
    run.font.name = FONT_DISPLAY if black else (FONT_BOLD if bold else FONT_BODY)
    run.font.size = Pt(size)
    run.font.bold = bold or black
    run.font.color.rgb = color


class CuikDeck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self._blank = self.prs.slide_layouts[6]
        self._n = 0

    # ── primitives ──────────────────────────────────────────────────
    def _slide(self, bg: RGBColor = WHITE):
        s = self.prs.slides.add_slide(self._blank)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = bg
        self._n += 1
        return s

    def _rect(self, s, x, y, w, h, color: RGBColor, radius: float | None = None, line: RGBColor | None = None):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = line
            shape.line.width = Pt(0.75)
        if radius:
            shape.adjustments[0] = radius
        shape.shadow.inherit = False
        return shape

    def _text(self, s, x, y, w, h, text: str, size: int, color=INK, bold=False, black=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing: float = 1.1):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        lines = text.split("\n") if isinstance(text, str) else list(text)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            r = p.add_run()
            r.text = line
            _font(r, size, color, bold, black)
        return tb

    def _bullets(self, s, x, y, w, h, items: Iterable[str], size: int = 18, color=INK):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            p.line_spacing = 1.15
            dot = p.add_run()
            dot.text = "•  "
            _font(dot, size, BLUE, bold=True)
            r = p.add_run()
            r.text = item
            _font(r, size, color)
        return tb

    def _footer(self, s, dark: bool = False):
        color = WHITE if dark else MUTED
        self._text(s, MARGIN, SLIDE_H - Inches(0.5), Inches(6), Inches(0.3), "Cuik · cuik.org", 11, color)
        self._text(s, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.5), Inches(1), Inches(0.3),
                   str(self._n), 11, color, align=PP_ALIGN.RIGHT)

    def _mark(self, s, x, y, h=Inches(0.45), on_blue: bool = False):
        s.shapes.add_picture(LOGO_MARK_ON_BLUE if on_blue else LOGO_MARK, x, y, height=h)

    def _asterisk(self, s, x, y, h=Inches(0.6), color: str = "orange"):
        s.shapes.add_picture(ASTERISK[color], x, y, height=h)

    def _title(self, s, title: str, dark: bool = False, asterisk: bool = True):
        color = WHITE if dark else INK
        self._text(s, MARGIN, Inches(0.6), SLIDE_W - 2 * MARGIN - Inches(1), Inches(1.1), title, 36, color, black=True)
        if asterisk:
            self._asterisk(s, SLIDE_W - MARGIN - Inches(0.55), Inches(0.62), Inches(0.5), "orange" if not dark else "orange")
        self._mark(s, MARGIN, SLIDE_H - Inches(1.05), Inches(0.42), on_blue=dark)

    # ── slide types ─────────────────────────────────────────────────
    def cover(self, title: str, subtitle: str = "", tag: str = ""):
        s = self._slide(BLUE)
        s.shapes.add_picture(LOGO_MARK_ON_BLUE, MARGIN, MARGIN, height=Inches(1.1))
        if tag:
            self._text(s, MARGIN, Inches(2.6), Inches(8), Inches(0.4), tag.upper(), 12, WHITE, bold=True)
        self._text(s, MARGIN, Inches(3.0), Inches(9.5), Inches(2.2), title, 48, WHITE, black=True, line_spacing=1.0)
        if subtitle:
            self._text(s, MARGIN, Inches(5.3), Inches(9), Inches(1), subtitle, 20, WHITE)
        self._asterisk(s, SLIDE_W - MARGIN - Inches(1.4), SLIDE_H - MARGIN - Inches(1.4), Inches(1.4), "orange")
        self._text(s, MARGIN, SLIDE_H - Inches(0.6), Inches(6), Inches(0.3), "cuik.org", 12, WHITE)
        return s

    def section(self, title: str, kicker: str = ""):
        s = self._slide(WHITE)
        self._rect(s, 0, 0, Inches(0.35), SLIDE_H, BLUE)
        if kicker:
            self._text(s, Inches(1.2), Inches(2.6), Inches(9), Inches(0.4), kicker.upper(), 12, MUTED, bold=True)
        self._text(s, Inches(1.2), Inches(3.0), Inches(10), Inches(2), title, 44, INK, black=True, line_spacing=1.0)
        self._asterisk(s, Inches(1.2), Inches(5.2), Inches(0.7), "orange")
        self._footer(s)
        return s

    def bullets(self, title: str, items: Sequence[str], note: str = ""):
        s = self._slide(WHITE)
        self._title(s, title)
        self._bullets(s, MARGIN, Inches(2.0), SLIDE_W - 2 * MARGIN, Inches(4.2), items, size=20)
        if note:
            self._text(s, MARGIN, Inches(6.2), SLIDE_W - 2 * MARGIN, Inches(0.5), note, 13, MUTED)
        self._footer(s)
        return s

    def text(self, title: str, body: str, size: int = 20):
        s = self._slide(WHITE)
        self._title(s, title)
        self._text(s, MARGIN, Inches(2.0), SLIDE_W - 2 * MARGIN, Inches(4.5), body, size, INK, line_spacing=1.25)
        self._footer(s)
        return s

    def two_columns(self, title: str, left: Sequence[str], right: Sequence[str],
                    left_title: str = "", right_title: str = "", highlight_right: bool = True):
        s = self._slide(WHITE)
        self._title(s, title)
        col_w = (SLIDE_W - 2 * MARGIN - Inches(0.4)) / 2
        y = Inches(1.9)
        h = Inches(4.6)
        # left card
        self._rect(s, MARGIN, y, col_w, h, SURFACE, radius=0.06, line=LINE)
        if left_title:
            self._text(s, MARGIN + Inches(0.4), y + Inches(0.35), col_w - Inches(0.8), Inches(0.5), left_title, 18, MUTED, bold=True)
        self._bullets(s, MARGIN + Inches(0.4), y + Inches(1.0), col_w - Inches(0.8), h - Inches(1.2), left, size=17)
        # right card
        rx = MARGIN + col_w + Inches(0.4)
        if highlight_right:
            self._rect(s, rx, y, col_w, h, BLUE, radius=0.06)
            if right_title:
                self._text(s, rx + Inches(0.4), y + Inches(0.35), col_w - Inches(0.8), Inches(0.5), right_title, 18, WHITE, bold=True)
            self._bullets(s, rx + Inches(0.4), y + Inches(1.0), col_w - Inches(0.8), h - Inches(1.2), right, size=17, color=WHITE)
        else:
            self._rect(s, rx, y, col_w, h, SURFACE, radius=0.06, line=LINE)
            if right_title:
                self._text(s, rx + Inches(0.4), y + Inches(0.35), col_w - Inches(0.8), Inches(0.5), right_title, 18, MUTED, bold=True)
            self._bullets(s, rx + Inches(0.4), y + Inches(1.0), col_w - Inches(0.8), h - Inches(1.2), right, size=17)
        self._footer(s)
        return s

    def stats(self, title: str, items: Sequence[tuple[str, str]], note: str = ""):
        """items: [(value, label), ...] up to 4."""
        s = self._slide(WHITE)
        self._title(s, title)
        n = max(1, min(4, len(items)))
        gap = Inches(0.35)
        w = (SLIDE_W - 2 * MARGIN - gap * (n - 1)) / n
        y = Inches(2.2)
        for i, (value, label) in enumerate(items[:4]):
            x = MARGIN + (w + gap) * i
            self._rect(s, x, y, w, Inches(2.6), SURFACE, radius=0.08, line=LINE)
            self._text(s, x, y + Inches(0.5), w, Inches(1.1), value, 44, BLUE, black=True, align=PP_ALIGN.CENTER)
            self._text(s, x + Inches(0.3), y + Inches(1.65), w - Inches(0.6), Inches(0.8), label, 14, MUTED, align=PP_ALIGN.CENTER)
        if note:
            self._text(s, MARGIN, Inches(5.4), SLIDE_W - 2 * MARGIN, Inches(0.8), note, 14, MUTED)
        self._footer(s)
        return s

    def table(self, title: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
              col_widths: Sequence[float] | None = None, note: str = ""):
        s = self._slide(WHITE)
        self._title(s, title)
        n_rows, n_cols = len(rows) + 1, len(headers)
        x, y = MARGIN, Inches(1.9)
        w = SLIDE_W - 2 * MARGIN
        row_h = Inches(0.42)
        shape = s.shapes.add_table(n_rows, n_cols, x, y, w, row_h * n_rows)
        tbl = shape.table
        if col_widths:
            total = sum(col_widths)
            for i, cw in enumerate(col_widths):
                tbl.columns[i].width = int(w * cw / total)
        for c, htxt in enumerate(headers):
            cell = tbl.cell(0, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = htxt
            _font(r, 13, WHITE, bold=True)
            cell.margin_left = cell.margin_right = Inches(0.12)
        for ri, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                cell = tbl.cell(ri, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = SURFACE if ri % 2 == 0 else WHITE
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                r = p.add_run()
                r.text = str(val)
                _font(r, 13, INK)
                if isinstance(val, (int, float)) or str(val).replace(".", "").replace(",", "").isdigit():
                    p.alignment = PP_ALIGN.RIGHT
                cell.margin_left = cell.margin_right = Inches(0.12)
        if note:
            self._text(s, MARGIN, y + row_h * n_rows + Inches(0.3), w, Inches(0.6), note, 13, MUTED)
        self._footer(s)
        return s

    def callout(self, title: str, quote: str, source: str = ""):
        """One big statement on a blue block (like the '¿Sabías que?' post)."""
        s = self._slide(WHITE)
        self._title(s, title, asterisk=False)
        self._rect(s, MARGIN, Inches(2.0), SLIDE_W - 2 * MARGIN, Inches(4.2), BLUE, radius=0.06)
        self._text(s, MARGIN + Inches(0.8), Inches(2.7), SLIDE_W - 2 * MARGIN - Inches(1.6), Inches(2.4), quote, 30, WHITE, black=True, line_spacing=1.05)
        if source:
            self._text(s, MARGIN + Inches(0.8), Inches(5.3), Inches(8), Inches(0.5), source, 14, WHITE)
        self._asterisk(s, SLIDE_W - MARGIN - Inches(1.1), Inches(2.3), Inches(0.7), "orange")
        self._footer(s)
        return s

    def image(self, title: str, path: str, caption: str = "", max_h: float = 4.6):
        s = self._slide(WHITE)
        self._title(s, title)
        pic = s.shapes.add_picture(path, MARGIN, Inches(1.9), height=Inches(max_h))
        if pic.width > SLIDE_W - 2 * MARGIN:
            ratio = (SLIDE_W - 2 * MARGIN) / pic.width
            pic.width = int(pic.width * ratio)
            pic.height = int(pic.height * ratio)
        pic.left = int((SLIDE_W - pic.width) / 2)
        if caption:
            self._text(s, MARGIN, Inches(6.6), SLIDE_W - 2 * MARGIN, Inches(0.4), caption, 12, MUTED, align=PP_ALIGN.CENTER)
        self._footer(s)
        return s

    def closing(self, title: str = "Gracias", contact: str = "francesco.leon@cuik.org · cuik.org"):
        s = self._slide(BLUE)
        s.shapes.add_picture(LOGO_MARK_ON_BLUE, MARGIN, MARGIN, height=Inches(1.1))
        self._text(s, MARGIN, Inches(3.0), Inches(10), Inches(1.4), title, 48, WHITE, black=True)
        self._text(s, MARGIN, Inches(4.5), Inches(10), Inches(0.6), contact, 18, WHITE)
        self._text(s, MARGIN, SLIDE_H - Inches(0.7), Inches(8), Inches(0.4), "Hecho con amor en Lima, Perú", 12, WHITE)
        self._asterisk(s, SLIDE_W - MARGIN - Inches(1.4), SLIDE_H - MARGIN - Inches(1.4), Inches(1.4), "orange")
        return s

    def save(self, path: str):
        self.prs.save(path)
        return path
