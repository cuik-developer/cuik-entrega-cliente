"""
Cuik-branded Word helper (python-docx).

    from cuik_docx import CuikDoc
    d = CuikDoc()
    d.cover("Guía de tu programa de puntos", "Preparado para Nombre del negocio", "Setiembre 2026")
    d.h1("Cómo suman puntos tus clientes")
    d.p("Texto de cuerpo...")
    d.bullets(["Uno", "Dos"])
    d.callout("Los puntos vencen el primer jueves de cada mes.")
    d.table(["Premio", "Puntos"], [["Café", "50"], ["Postre", "80"]])
    d.kv([("Regla", "Cada compra vence a los 30 días"), ("Aviso", "2 días antes, 10:00")])
    d.save("salida.docx")
"""

from __future__ import annotations

import os
from typing import Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(os.path.dirname(HERE), "assets")
LOGO_MARK = os.path.join(ASSETS, "logo", "logo-mark.jpeg")
WORDMARK = os.path.join(ASSETS, "logo", "wordmark-blue-on-white.png")
ASTERISK_ORANGE = os.path.join(ASSETS, "logo", "asterisk-orange.png")

BLUE = RGBColor(0x0E, 0x70, 0xDB)
ORANGE = RGBColor(0xFF, 0x48, 0x10)
INK = RGBColor(0x23, 0x1F, 0x20)
MUTED = RGBColor(0x6B, 0x72, 0x80)
BLUE_HEX = "0E70DB"
SURFACE_HEX = "F8F8F8"
LINE_HEX = "E5E7EB"

FONT = "Poppins"
FONT_DISPLAY = "Poppins Black"


def _shade(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _borders(table, color: str = LINE_HEX, size: int = 4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _run(p, text: str, size: float, color: RGBColor = INK, bold: bool = False, display: bool = False):
    r = p.add_run(text)
    r.font.name = FONT_DISPLAY if display else FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_DISPLAY if display else FONT)
    r.font.size = Pt(size)
    r.font.bold = bold or display
    r.font.color.rgb = color
    return r


class CuikDoc:
    def __init__(self):
        self.doc = Document()
        for section in self.doc.sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.2)
            section.right_margin = Cm(2.2)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(11)
        st.font.color.rgb = INK
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.25
        self._footer()

    def _footer(self):
        f = self.doc.sections[0].footer
        p = f.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _run(p, "Cuik · cuik.org", 9, MUTED)

    # ── blocks ──────────────────────────────────────────────────────
    def cover(self, title: str, subtitle: str = "", tag: str = ""):
        self.doc.add_picture(LOGO_MARK, height=Cm(1.6))
        self.doc.add_paragraph()
        if tag:
            p = self.doc.add_paragraph()
            _run(p, tag.upper(), 9, MUTED, bold=True)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _run(p, title, 30, BLUE, display=True)
        if subtitle:
            p = self.doc.add_paragraph()
            _run(p, subtitle, 13, MUTED)
        p = self.doc.add_paragraph()
        p.add_run().add_picture(ASTERISK_ORANGE, height=Cm(0.9))
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        _run(p, text, 18, INK, display=True)

    def h2(self, text: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        _run(p, text, 13, BLUE, bold=True)

    def p(self, text: str, muted: bool = False):
        p = self.doc.add_paragraph()
        _run(p, text, 11, MUTED if muted else INK)
        return p

    def bullets(self, items: Sequence[str]):
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            _run(p, it, 11)

    def numbered(self, items: Sequence[str]):
        for it in items:
            p = self.doc.add_paragraph(style="List Number")
            _run(p, it, 11)

    def callout(self, text: str, label: str = ""):
        """Blue block with white text: one key statement."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        _shade(cell, BLUE_HEX)
        cell.paragraphs[0].paragraph_format.space_before = Pt(8)
        if label:
            _run(cell.paragraphs[0], label.upper(), 8, RGBColor(0xDB, 0xEA, 0xFE), bold=True)
            p = cell.add_paragraph()
        else:
            p = cell.paragraphs[0]
        _run(p, text, 14, RGBColor(0xFF, 0xFF, 0xFF), display=True)
        p.paragraph_format.space_after = Pt(8)
        self.doc.add_paragraph()

    def table(self, headers: Sequence[str], rows: Sequence[Sequence[str]], widths_cm: Sequence[float] | None = None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            _shade(c, BLUE_HEX)
            c.paragraphs[0].text = ""
            _run(c.paragraphs[0], h, 10, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].paragraphs[0].text = ""
                _run(cells[i].paragraphs[0], str(v), 10)
                if ri % 2 == 1:
                    _shade(cells[i], SURFACE_HEX)
                if str(v).replace(".", "").replace(",", "").isdigit():
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths_cm:
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    def kv(self, pairs: Sequence[tuple[str, str]]):
        """Two-column definition list (label muted, value ink)."""
        t = self.doc.add_table(rows=0, cols=2)
        _borders(t)
        for k, v in pairs:
            cells = t.add_row().cells
            _shade(cells[0], SURFACE_HEX)
            cells[0].paragraphs[0].text = ""
            _run(cells[0].paragraphs[0], k, 10, MUTED, bold=True)
            cells[1].paragraphs[0].text = ""
            _run(cells[1].paragraphs[0], v, 10)
            cells[0].width = Cm(5.5)
            cells[1].width = Cm(11)
        self.doc.add_paragraph()
        return t

    def stat_row(self, items: Sequence[tuple[str, str]]):
        """Up to 4 big numbers with labels."""
        t = self.doc.add_table(rows=2, cols=len(items))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, LINE_HEX, 4)
        for i, (val, label) in enumerate(items):
            c0 = t.cell(0, i)
            c1 = t.cell(1, i)
            _shade(c0, SURFACE_HEX)
            _shade(c1, SURFACE_HEX)
            c0.paragraphs[0].text = ""
            c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(c0.paragraphs[0], val, 24, BLUE, display=True)
            c1.paragraphs[0].text = ""
            c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(c1.paragraphs[0], label, 9, MUTED)
        self.doc.add_paragraph()

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path: str):
        self.doc.save(path)
        return path
